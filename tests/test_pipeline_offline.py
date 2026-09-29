"""离线端到端：用合成夹具走完整链路（零网络、零费用）。

覆盖：附件分组 → XML/PDF 抽取 → 抬头/用途归类 → 按发票号码去重 → 行程聚类 →
      目录与命名落盘 → 审计 JSON / Excel / 分类合并 PDF → 幂等闸门。
"""

from __future__ import annotations

import json
import sys
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from fapiao import archive, pipeline  # noqa: E402
from tests import fixtures, support  # noqa: E402

RUN_DATE = date(2026, 10, 20)


def _attachment(uid, index, filename, path, kind, role, size=100):
    return {
        "mail_uid": uid, "mailbox": "INBOX", "index": index, "filename": filename,
        "content_type": "", "size": size, "path": str(path), "kind": kind, "role": role,
    }


def build_staging(base: Path) -> Path:
    staging = base / "staging"
    staging.mkdir(parents=True, exist_ok=True)
    mails = []

    # 邮件 1：一张杭州打车发票（PDF + XML + OFD + 行程单 PDF）——4 个附件只算 1 张发票。
    p1 = staging / "001_01_高德打车电子发票.pdf"
    fixtures.make_pdf(p1, fixtures.VAT_TAXI_INVOICE_TEXT)
    x1 = staging / "001_02_高德打车电子发票.xml"
    x1.write_text(fixtures.TAXI_INVOICE_XML, encoding="utf-8")
    o1 = staging / "001_03_高德打车电子发票.ofd"
    o1.write_bytes(b"OFD-DUMMY")
    i1 = staging / "001_04_高德打车电子行程单.pdf"
    fixtures.make_pdf(i1, fixtures.ITINERARY_TEXT)
    mails.append({
        "mail": {"uid": "101", "mailbox": "INBOX", "subject": "高德打车电子发票",
                 "sender": "itinerary@example.com", "date": "2026-09-18T05:39:09+00:00",
                 "message_id": "<1@example.com>", "text": "共计1单行程，合计22.23元"},
        "attachments": [
            _attachment("101", 1, "高德打车电子发票.pdf", p1, "invoice", "pdf"),
            _attachment("101", 2, "高德打车电子发票.xml", x1, "invoice", "xml"),
            _attachment("101", 3, "高德打车电子发票.ofd", o1, "invoice", "ofd"),
            _attachment("101", 4, "高德打车电子行程单.pdf", i1, "itinerary", "pdf"),
        ],
    })

    # 邮件 2：杭州东→上海南 高铁票（只带 PDF，验证铁路分支与行程闭环）。
    p2 = staging / "002_01_电子发票-RAIL001.pdf"
    fixtures.make_pdf(p2, fixtures.RAIL_INVOICE_TEXT)
    mails.append({
        "mail": {"uid": "102", "mailbox": "INBOX", "subject": "火车票电子发票",
                 "sender": "train@example.com", "date": "2026-09-18T03:57:56+00:00",
                 "message_id": "<2@example.com>", "text": ""},
        "attachments": [_attachment("102", 1, "电子发票-RAIL001.pdf", p2, "invoice", "pdf")],
    })

    # 邮件 3：与邮件 1 同一张发票（同发票号码）——必须被去重，不得重复计入金额。
    p3 = staging / "003_01_高德打车电子发票.pdf"
    fixtures.make_pdf(p3, fixtures.VAT_TAXI_INVOICE_TEXT)
    mails.append({
        "mail": {"uid": "103", "mailbox": "INBOX", "subject": "高德打车电子发票（重复投递）",
                 "sender": "itinerary@example.com", "date": "2026-09-19T01:00:00+00:00",
                 "message_id": "<3@example.com>", "text": ""},
        "attachments": [_attachment("103", 1, "高德打车电子发票.pdf", p3, "invoice", "pdf")],
    })

    # 邮件 4：只有 XML、没有 PDF——应判"缺少PDF版发票文件"并进存疑。
    x4 = staging / "004_01_住宿电子发票.xml"
    x4.write_text(fixtures.HOTEL_INVOICE_XML_ONLY, encoding="utf-8")
    mails.append({
        "mail": {"uid": "104", "mailbox": "INBOX", "subject": "酒店订单电子发票",
                 "sender": "hotel@example.com", "date": "2026-09-17T08:54:07+00:00",
                 "message_id": "<4@example.com>", "text": ""},
        "attachments": [_attachment("104", 1, "住宿电子发票.xml", x4, "invoice", "xml")],
    })

    manifest = {
        "mailbox": "INBOX",
        "window": {"since": "2026-09-21", "before": "2026-10-21"},
        "mail_count": len(mails),
        "mails": mails,
    }
    (staging / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8", newline=""
    )
    return staging


class TestOfflinePipeline(unittest.TestCase):
    def setUp(self):
        self.base = support.make_temp_dir("pipeline")
        self.staging = build_staging(self.base)
        self.out_root = self.base / "归档"
        self.cfg = fixtures.test_config(self.out_root)

    def tearDown(self):
        support.clean_temp_dir(self.base)

    def _run(self, **kwargs):
        return pipeline.run(
            self.cfg, RUN_DATE, replay_dir=self.staging, log=lambda *_a, **_k: None, **kwargs
        )

    def test_end_to_end_archive(self):
        result = self._run()
        self.assertIsNotNone(result.run_dir)
        run_dir = result.run_dir
        self.assertEqual(run_dir.parent, self.out_root)
        self.assertTrue(run_dir.name.startswith("20261020月度发票归集("))

        summary = result.summary
        self.assertEqual(summary["archive_invoice_count"], 4)
        self.assertEqual(summary["included_invoice_count"], 2)
        self.assertEqual(summary["included_total"], 101.23)
        self.assertEqual(summary["excluded_invoice_count"], 2)
        self.assertEqual(summary["category_totals"]["测试大学杭州市内打车票"], 22.23)
        self.assertEqual(summary["category_totals"]["测试大学差旅费"], 79.00)

        # 4 个附件只算 1 张发票；OFD 被忽略并留痕。
        self.assertEqual(len(summary["skipped"]["duplicates"]), 1)
        self.assertTrue(any("OFD" in item["reason"] for item in summary["skipped"]["attachments"]))

        # 归档产物齐全
        for name in ("00_审计摘要.json", "00_发票审计明细及汇总.xlsx", "00_审计明细.json"):
            self.assertTrue((run_dir / name).is_file(), name)
        merged_dir = run_dir / "00_汇总发票"
        self.assertTrue((merged_dir / "合并清单.json").is_file())
        merged_files = sorted(p.name for p in merged_dir.glob("*.pdf"))
        self.assertEqual(
            merged_files,
            sorted(["测试大学差旅费-79.00元.pdf", "测试大学杭州市内打车票-22.23元.pdf"]),
        )

        # 明细目录按 config 里的抬头顺序编号；行程文件夹命名复刻参考案例
        trip = run_dir / "01_测试大学明细" / "01_差旅费" / "01_2026-07-18_至_2026-07-18_上海差旅_杭州-上海-杭州"
        self.assertTrue(trip.is_dir(), trip)
        rail_files = list(trip.glob("*_发票.pdf"))
        self.assertEqual(len(rail_files), 1)
        self.assertTrue(rail_files[0].name.endswith("_79.00元_发票649100_发票.pdf"), rail_files[0].name)

        taxi_dir = run_dir / "01_测试大学明细" / "03_杭州市内打车票"
        taxi_files = list(taxi_dir.glob("*_发票.pdf"))
        self.assertEqual(len(taxi_files), 1)
        self.assertEqual(len(list(taxi_dir.glob("*_行程单.pdf"))), 1)
        self.assertTrue(taxi_files[0].name.endswith("_22.23元_发票508042_发票.pdf"), taxi_files[0].name)

        # _原始数据 里存了 XML
        self.assertEqual(len(list((run_dir / "_原始数据").glob("*.xml"))), 2)

        # 存疑目录：重复 + 缺 PDF
        quarantine = run_dir / "99_存疑待人工审计"
        self.assertTrue((quarantine / "重复发票不重复计入").is_dir())
        self.assertTrue((quarantine / "缺少PDF版发票文件").is_dir())

        # 三处金额一致由 verify_run 保证；这里再独立复核一次
        self.assertEqual(
            round(sum(v for v in summary["category_totals"].values()), 2),
            summary["included_total"],
        )
        self.assertEqual(
            round(sum(m["amount"] for m in summary["分类汇总PDF"]), 2),
            summary["included_total"],
        )

    def test_rerun_refuses_to_overwrite(self):
        self._run()
        with self.assertRaises(FileExistsError):
            self._run()

    def test_force_rebuilds_identically(self):
        first = self._run()
        second = self._run(force=True)
        self.assertEqual(first.summary["included_total"], second.summary["included_total"])
        self.assertEqual(first.summary["archive_invoice_count"], second.summary["archive_invoice_count"])

    def test_dry_run_writes_nothing(self):
        result = self._run(dry_run=True)
        self.assertIsNone(result.run_dir)
        self.assertEqual(result.summary["included_total"], 101.23)
        self.assertFalse(self.out_root.exists())

    def test_excel_has_both_sheets_and_frozen_headers(self):
        result = self._run()
        from openpyxl import load_workbook

        workbook = load_workbook(result.run_dir / "00_发票审计明细及汇总.xlsx")
        self.assertEqual(workbook.sheetnames, ["汇总", "审计明细"])
        audit = workbook["审计明细"]
        header = [audit.cell(row=4, column=i).value for i in range(1, 14)]
        self.assertEqual(header, [
            "序号", "抬头", "用途", "行程日期", "地点或路线", "票据类型", "金额", "发票号码",
            "归档状态", "金额口径", "归档发票", "计入总额", "分类",
        ])
        self.assertEqual(audit.freeze_panes, "D5")
        self.assertEqual(audit.max_row, 4 + result.summary["archive_invoice_count"])


class TestOwnerDir(unittest.TestCase):
    def test_config_driven_numbering(self):
        cfg = {"classify": {"buyers": [{"name": "浙江大学"}, {"name": "芯智辰光"}]}}
        self.assertEqual(archive.owner_dir("浙江大学", cfg), "01_浙江大学明细")
        self.assertEqual(archive.owner_dir("芯智辰光", cfg), "02_芯智辰光明细")
        self.assertIsNone(archive.owner_dir("未知公司", cfg))


if __name__ == "__main__":
    unittest.main()
